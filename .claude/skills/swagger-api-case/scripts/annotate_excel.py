"""
annotate_excel.py — 给 swagger sheet 的「备注」列批量写说明（零流量 / 异步导出 / 仅生成未执行 等）
用法: python annotate_excel.py <xlsx> <sheet> <notes.json>
notes.json: [{"method":"POST","path":"/api/xxx","note":"零流量(365d 0次)，暂不覆盖","scene":"可选，写入场景覆盖列"}]
🔴 串行使用，禁止多个进程同时读写同一份 xlsx
"""
import json, sys, openpyxl
from openpyxl.styles import Alignment

xlsx, sheet, notes_p = sys.argv[1], sys.argv[2], sys.argv[3]
notes = json.load(open(notes_p, encoding='utf-8'))
wb = openpyxl.load_workbook(xlsx)
ws = wb[sheet]
hdr = {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value}
pc, mc = hdr['接口路径'], hdr['Method']
nc = hdr.get('备注') or ws.max_column + 1
if '备注' not in hdr:
    ws.cell(1, nc, '备注')
sc = hdr.get('场景覆盖')
idx = {(str(ws.cell(r, mc).value or '').upper(), ws.cell(r, pc).value): r for r in range(2, ws.max_row + 1)}
n = 0
for it in notes:
    r = idx.get((it['method'].upper(), it['path']))
    if not r:
        print('  not found:', it['method'], it['path'])
        continue
    ws.cell(r, nc, it['note']).alignment = Alignment(wrap_text=True, vertical='top')
    n += 1
    if it.get('scene') and sc:
        ws.cell(r, sc, it['scene']).alignment = Alignment(wrap_text=True, vertical='top')
wb.save(xlsx)
print(f"annotated {n}/{len(notes)} rows in {sheet}")
