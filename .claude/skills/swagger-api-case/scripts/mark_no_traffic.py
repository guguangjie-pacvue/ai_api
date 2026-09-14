#!/usr/bin/env python3
"""
mark_no_traffic.py — Phase 5 补充工具：标注"ES 无流量"的接口（不生成 cases.json/report.json）。

用途：某些接口在目标平台下 ES 近 90 天无真实流量（业务性零结果，且该接口本身按平台可归因），
按 SKILL.md 红线：严禁生成 cases.json / 执行 run_cases.py，唯一动作是在 Excel 里标注：
  - RuleApi sheet 对应行的「场景覆盖」列 = "ES无流量"
  - 若整个模块的接口全部是 ES 无流量（模块内 0 条真实 case），则「模块汇总」行的
    Case数=0，通过率留空，接口覆盖率="0% (0/api_total)"
  - 若模块内部分接口有真实 case（已由 update_excel.py 正常流程写过模块汇总行），
    则不重复改写模块汇总行，只补写 RuleApi sheet 里剩余无流量接口的场景覆盖列。

用法：
  python3 mark_no_traffic.py \
    --excel single-api/services/rule-api/swagger_modules.xlsx \
    --sheet-name RuleApi \
    --swagger-title samsclub \
    --module feature-usage-controller \
    --env us \
    --endpoints single-api/endpoints-RuleApi.json \
    --paths "POST:/featureHealth/monthlyReport,POST:/featureHealth/topNClients" \
    --whole-module-empty
"""
import argparse
import json
import shutil

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def _thin_border():
    s = Side(style='thin', color='BFBFBF')
    return Border(left=s, right=s, top=s, bottom=s)


def _find_or_create_col(ws, col_name, header_font=None, header_fill=None, width=None):
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value == col_name:
            return c
    col_idx = ws.max_column + 1
    cell = ws.cell(1, col_idx, col_name)
    if header_font:
        cell.font = header_font
    if header_fill:
        cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = _thin_border()
    if width:
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    return col_idx


def mark_scenario_cells(wb, sheet_name, env, swagger_title, module, method_paths, label='ES无流量'):
    if sheet_name not in wb.sheetnames:
        print(f'[mark_no_traffic] sheet "{sheet_name}" not found, skip')
        return 0
    ws = wb[sheet_name]
    bdr = _thin_border()
    h_font = Font(bold=True, color='FFFFFF')
    h_fill = PatternFill('solid', fgColor='2E5F8A')

    env_col = next((c for c in range(1, ws.max_column + 1) if '环境' in str(ws.cell(1, c).value or '')), None)
    path_col = next((c for c in range(1, ws.max_column + 1)
                      if '路径' in str(ws.cell(1, c).value or '') or 'path' in str(ws.cell(1, c).value or '').lower()), None)
    method_col = next((c for c in range(1, ws.max_column + 1)
                        if str(ws.cell(1, c).value or '').strip().lower() in ('method', '方法', 'http方法')), None)
    platform_col = next((c for c in range(1, ws.max_column + 1) if 'Platform' in str(ws.cell(1, c).value or '')), None)
    module_col = next((c for c in range(1, ws.max_column + 1) if 'Tag' in str(ws.cell(1, c).value or '')), None)
    if path_col is None:
        print(f'[mark_no_traffic] path column not found in {sheet_name}')
        return 0

    scene_col = _find_or_create_col(ws, '场景覆盖', h_font, h_fill, width=72)

    updated = 0
    for r in range(2, ws.max_row + 1):
        if env_col and ws.cell(r, env_col).value not in (env, None, ''):
            continue
        if platform_col and ws.cell(r, platform_col).value != swagger_title:
            continue
        if module_col and ws.cell(r, module_col).value != module:
            continue
        row_path = ws.cell(r, path_col).value
        row_method = (ws.cell(r, method_col).value or '').strip().upper() if method_col else ''
        if not row_path:
            continue
        if (row_method, row_path) in method_paths:
            cell = ws.cell(r, scene_col, label)
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            cell.border = bdr
            updated += 1
    print(f'[mark_no_traffic] {sheet_name}[{env}] {swagger_title}/{module}: marked {updated} rows as "{label}"')
    return updated


def mark_summary_row(wb, module, env, swagger_title, api_total):
    ws = wb['模块汇总']
    col_idx = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value or ''
        if '环境' in v:
            col_idx['env'] = c
        if 'Swagger' in v:
            col_idx['swagger'] = c
        if 'Platform' in v:
            col_idx['platform'] = c
        if 'Tag' in v:
            col_idx['module'] = c
    summary_font = Font(bold=True)
    summary_fill = PatternFill('solid', fgColor='D9E1F2')
    for name in ('Case数', '通过率', '接口覆盖率'):
        col_idx[name] = _find_or_create_col(ws, name, summary_font, summary_fill)

    has_swagger = 'swagger' in col_idx
    has_platform = 'platform' in col_idx

    for r in range(2, ws.max_row + 1):
        env_val = ws.cell(r, col_idx['env']).value if 'env' in col_idx else None
        env_ok = env_val in (env, None, '')
        mod_ok = 'module' in col_idx and ws.cell(r, col_idx['module']).value == module
        if has_swagger:
            key_ok = ws.cell(r, col_idx['swagger']).value == swagger_title
        elif has_platform:
            key_ok = ws.cell(r, col_idx['platform']).value == swagger_title
        else:
            key_ok = True
        if env_ok and key_ok and mod_ok:
            ws.cell(r, col_idx['Case数'], 0).alignment = Alignment(horizontal='center')
            pass_cell = ws.cell(r, col_idx['通过率'])
            pass_cell.value = None  # explicit clear: ws.cell(..., None) does NOT clear in openpyxl
            pass_cell.alignment = Alignment(horizontal='center')
            coverage = f'0% (0/{api_total})' if api_total else 'N/A'
            ws.cell(r, col_idx['接口覆盖率'], coverage).alignment = Alignment(horizontal='center')
            print(f'[mark_no_traffic][模块汇总] {swagger_title}/{module}[{env}]: case=0, pass=<blank>, coverage={coverage}')
            return
    print(f'[mark_no_traffic][模块汇总] WARNING: row not found for {swagger_title}/{module}[{env}]')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', required=True)
    ap.add_argument('--sheet-name', required=True, dest='sheet_name')
    ap.add_argument('--swagger-title', required=True, dest='swagger_title')
    ap.add_argument('--module', required=True)
    ap.add_argument('--env', default='us', choices=['us', 'cn', 'eu'])
    ap.add_argument('--endpoints', required=True)
    ap.add_argument('--paths', required=True,
                    help='逗号分隔的 METHOD:/path 列表，标记这些接口为 ES无流量')
    ap.add_argument('--whole-module-empty', action='store_true',
                    help='该模块在本平台下所有接口均无流量（模块汇总也写 Case数=0）')
    ap.add_argument('--label', default='ES无流量',
                    help='写入场景覆盖列的文本，默认"ES无流量"；风险跳过场景可传"高风险跳过（未执行，见 RISK_NOTES.md）"')
    args = ap.parse_args()

    method_paths = set()
    for item in args.paths.split(','):
        item = item.strip()
        if not item:
            continue
        method, path = item.split(':', 1)
        method_paths.add((method.strip().upper(), path.strip()))

    with open(args.endpoints, encoding='utf-8-sig') as f:
        endpoints = json.load(f)
    api_total = len([e for e in endpoints if e.get('tag') == args.module])

    wb = openpyxl.load_workbook(args.excel)
    mark_scenario_cells(wb, args.sheet_name, args.env, args.swagger_title, args.module, method_paths, label=args.label)
    if args.whole_module_empty:
        mark_summary_row(wb, args.module, args.env, args.swagger_title, api_total)

    tmp = args.excel.replace('.xlsx', '_tmp.xlsx')
    wb.save(tmp)
    try:
        shutil.move(tmp, args.excel)
        print('Excel saved:', args.excel)
    except PermissionError:
        print('Excel 被占用（请关闭后手动用 tmp 覆盖）。tmp 已保存:', tmp)


if __name__ == '__main__':
    main()
