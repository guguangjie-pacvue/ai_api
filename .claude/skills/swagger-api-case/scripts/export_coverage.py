"""
export_coverage.py — 把每份 swagger_modules.xlsx 转成同目录下的 coverage.json，一对一，不汇总。

现状：mainapi / walmart-api / micro-api / rule-api 各自维护一份 swagger_modules.xlsx。
这个脚本不重新计算覆盖率（Excel 仍是唯一数据源，由 update_excel.py 维护），只是原样
读出每份 Excel 的"模块汇总" sheet + 其余的接口明细 sheet，转成 JSON，写在该 Excel 同目录下。

Usage:
  python scripts/export_coverage.py
  # 每份 single-api/**/swagger_modules.xlsx 旁边生成一份 coverage.json
"""
import glob
import json
import os

import openpyxl

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
SINGLE_API = os.path.join(REPO_ROOT, 'single-api')

SUMMARY_SHEET = '模块汇总'


def _sheet_to_rows(ws):
    """把一个 sheet 转成 list[dict]，key 用表头原文。
    第一列（环境）在 Excel 里常常只在块首行填值、后续行留空表示"继承上一行"，
    这里做前向填充，否则 JSON 里一半的行会丢失环境信息。"""
    rows_iter = ws.iter_rows(values_only=True)
    headers = list(next(rows_iter))
    rows = []
    carry_first_col = None
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        row = list(raw)
        if row[0] not in (None, ''):
            carry_first_col = row[0]
        else:
            row[0] = carry_first_col
        rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return rows


def convert_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    summary = _sheet_to_rows(wb[SUMMARY_SHEET]) if SUMMARY_SHEET in wb.sheetnames else []
    details = {
        name: _sheet_to_rows(wb[name])
        for name in wb.sheetnames
        if name != SUMMARY_SHEET
    }
    return summary, details


def main():
    paths = sorted(glob.glob(os.path.join(SINGLE_API, '**', 'swagger_modules.xlsx'), recursive=True))
    for path in paths:
        rel = os.path.relpath(path, REPO_ROOT).replace('\\', '/')
        # 服务名取 swagger_modules.xlsx 所在的目录名（不用路径第二段，rule-api/micro-api
        # 嵌套在 services/ 、micro-server/ 下，第二段会拿到中间目录名而不是服务名）
        service = rel.split('/')[-2]
        summary, details = convert_workbook(path)
        data = {
            'schema_version': 2,
            'service': service,
            'excel_path': rel,
            'summary': summary,
            'details': details,
        }

        out_json = os.path.join(os.path.dirname(path), 'coverage.json')
        with open(out_json, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        out_rel = os.path.relpath(out_json, REPO_ROOT).replace('\\', '/')
        print(f'[{service}] {rel} -> {out_rel}  (模块汇总 {len(summary)} 行, 明细 sheet {list(details.keys())})')


if __name__ == '__main__':
    main()
